"""
Document Parsing Module for PredictX RAG Pipeline.
Extracts text and page-level metadata from PDF, DOCX, TXT, and CSV documents.
"""
import csv
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any

import fitz  # PyMuPDF
from docx import Document as DocxDocument

from ingestion.document_validator import compute_file_hash, compute_text_hash
from ingestion.preprocessing import clean_text, is_meaningful_text, extract_headings

class ParsedPage:
    def __init__(self, text: str, page_number: int, section_name: str = "", metadata: Dict[str, Any] = None):
        self.text = text
        self.page_number = page_number
        self.section_name = section_name
        self.metadata = metadata or {}

class ParsedDocument:
    def __init__(self, document_id: str, document_name: str, document_type: str, category: str, pages: List[ParsedPage], file_hash: str):
        self.document_id = document_id
        self.document_name = document_name
        self.document_type = document_type
        self.category = category
        self.pages = pages
        self.file_hash = file_hash
        self.total_pages = len(pages)

def detect_machine_context(text: str) -> Dict[str, str]:
    """
    Infers machine_type and primary component from document text keywords.
    """
    text_lower = text.lower()
    machine_type = "GENERAL_INDUSTRIAL"
    component = "GENERAL"

    if "turbofan" in text_lower or "c-mapss" in text_lower or "jet engine" in text_lower:
        machine_type = "TURBINE"
    elif "cnc" in text_lower or "milling" in text_lower or "lathe" in text_lower:
        machine_type = "CNC_MILL"
    elif "press" in text_lower or "hydraulic" in text_lower:
        machine_type = "HYDRAULIC_PRESS"
    elif "robotic" in text_lower or "arm" in text_lower:
        machine_type = "ROBOTIC_ARM"

    if "compressor" in text_lower or "hpc" in text_lower or "lpc" in text_lower:
        component = "COMPRESSOR"
    elif "turbine" in text_lower or "hpt" in text_lower or "lpt" in text_lower:
        component = "TURBINE"
    elif "bearing" in text_lower or "raceway" in text_lower or "bpfo" in text_lower:
        component = "BEARING"
    elif "tool" in text_lower or "insert" in text_lower or "wear" in text_lower:
        component = "CUTTING_TOOL"
    elif "spindle" in text_lower:
        component = "SPINDLE"
    elif "cooling" in text_lower or "heat exchanger" in text_lower or "overheating" in text_lower:
        component = "COOLING_SYSTEM"

    return {"machine_type": machine_type, "component": component}

def parse_pdf(file_path: Path, category: str) -> ParsedDocument:
    doc = fitz.open(str(file_path))
    file_hash = compute_file_hash(file_path)
    document_id = f"DOC-{file_hash[:10].upper()}"
    document_name = file_path.name
    
    pages = []
    current_section = category.replace("_", " ").title()

    for idx, page in enumerate(doc, start=1):
        raw_text = page.get_text("text")
        cleaned = clean_text(raw_text)
        
        if not is_meaningful_text(cleaned):
            continue

        headings = extract_headings(cleaned)
        if headings:
            current_section = headings[0]

        context = detect_machine_context(cleaned)
        metadata = {
            "document_id": document_id,
            "document_name": document_name,
            "document_type": "PDF",
            "category": category,
            "page_number": idx,
            "section_name": current_section,
            "machine_type": context["machine_type"],
            "component": context["component"],
            "uploaded_at": datetime.utcnow().isoformat(),
            "content_hash": compute_text_hash(cleaned)
        }
        pages.append(ParsedPage(text=cleaned, page_number=idx, section_name=current_section, metadata=metadata))

    doc.close()
    return ParsedDocument(document_id, document_name, "PDF", category, pages, file_hash)

def parse_docx(file_path: Path, category: str) -> ParsedDocument:
    doc = DocxDocument(str(file_path))
    file_hash = compute_file_hash(file_path)
    document_id = f"DOC-{file_hash[:10].upper()}"
    document_name = file_path.name

    # Group paragraphs under headings
    pages = []
    current_section = category.replace("_", " ").title()
    current_text_blocks = []
    page_counter = 1

    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        
        if p.style.name.startswith("Heading") or (len(txt) < 80 and txt.isupper()):
            if current_text_blocks:
                full_page_text = clean_text("\n\n".join(current_text_blocks))
                if is_meaningful_text(full_page_text):
                    context = detect_machine_context(full_page_text)
                    meta = {
                        "document_id": document_id,
                        "document_name": document_name,
                        "document_type": "DOCX",
                        "category": category,
                        "page_number": page_counter,
                        "section_name": current_section,
                        "machine_type": context["machine_type"],
                        "component": context["component"],
                        "uploaded_at": datetime.utcnow().isoformat(),
                        "content_hash": compute_text_hash(full_page_text)
                    }
                    pages.append(ParsedPage(full_page_text, page_counter, current_section, meta))
                    page_counter += 1
                current_text_blocks = []
            current_section = txt
        else:
            current_text_blocks.append(txt)

    # Process remaining paragraphs
    if current_text_blocks:
        full_page_text = clean_text("\n\n".join(current_text_blocks))
        if is_meaningful_text(full_page_text):
            context = detect_machine_context(full_page_text)
            meta = {
                "document_id": document_id,
                "document_name": document_name,
                "document_type": "DOCX",
                "category": category,
                "page_number": page_counter,
                "section_name": current_section,
                "machine_type": context["machine_type"],
                "component": context["component"],
                "uploaded_at": datetime.utcnow().isoformat(),
                "content_hash": compute_text_hash(full_page_text)
            }
            pages.append(ParsedPage(full_page_text, page_counter, current_section, meta))

    return ParsedDocument(document_id, document_name, "DOCX", category, pages, file_hash)

def parse_txt(file_path: Path, category: str) -> ParsedDocument:
    file_hash = compute_file_hash(file_path)
    document_id = f"DOC-{file_hash[:10].upper()}"
    document_name = file_path.name

    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    raw_text = ""
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                raw_text = f.read()
            break
        except UnicodeDecodeError:
            continue

    cleaned = clean_text(raw_text)
    pages = []
    
    # Split text into logical sections based on major headers or page breaks
    sections = cleaned.split("\n\n---")
    if len(sections) == 1:
        # Check for numeric section dividers like "1. ", "2. ", "Section "
        import re
        sections = re.split(r"\n(?=(?:Section\s+\d+|Chapter\s+\d+|\d+\.\s+[A-Z]))", cleaned)

    for idx, sec in enumerate(sections, start=1):
        sec_clean = clean_text(sec)
        if not is_meaningful_text(sec_clean):
            continue

        headings = extract_headings(sec_clean)
        section_name = headings[0] if headings else f"Section {idx}"
        context = detect_machine_context(sec_clean)
        
        meta = {
            "document_id": document_id,
            "document_name": document_name,
            "document_type": "TXT",
            "category": category,
            "page_number": idx,
            "section_name": section_name,
            "machine_type": context["machine_type"],
            "component": context["component"],
            "uploaded_at": datetime.utcnow().isoformat(),
            "content_hash": compute_text_hash(sec_clean)
        }
        pages.append(ParsedPage(sec_clean, idx, section_name, meta))

    return ParsedDocument(document_id, document_name, "TXT", category, pages, file_hash)

def parse_csv(file_path: Path, category: str) -> ParsedDocument:
    file_hash = compute_file_hash(file_path)
    document_id = f"DOC-{file_hash[:10].upper()}"
    document_name = file_path.name

    pages = []
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        reader = csv.DictReader(f)
        for idx, row in enumerate(reader, start=1):
            # Convert structured table record to high-density semantic natural language
            rec_id = row.get("record_id", f"REC-{idx}")
            machine_id = row.get("machine_id", "UNKNOWN")
            m_type = row.get("machine_type", "INDUSTRIAL_MACHINE")
            failure_mode = row.get("failure_mode", "General Failure")
            sensor = row.get("dominant_sensor", "N/A")
            val = row.get("sensor_value_at_failure", "N/A")
            thresh = row.get("threshold_limit", "N/A")
            action = row.get("remedial_action", "Inspect and service")
            downtime = row.get("downtime_hours", "N/A")
            cost = row.get("cost_usd", "N/A")

            semantic_record = (
                f"Historical Maintenance Incident Record [{rec_id}]:\n"
                f"- Machine ID: {machine_id} (Type: {m_type})\n"
                f"- Diagnosed Failure Mode: {failure_mode}\n"
                f"- Dominant Sensor Telemetry at Incident: {sensor} = {val} (Threshold: {thresh})\n"
                f"- Corrective Action Executed: {action}\n"
                f"- Operational Impact: {downtime} hours downtime, total repair cost ${cost} USD."
            )

            meta = {
                "document_id": document_id,
                "document_name": document_name,
                "document_type": "CSV",
                "category": category,
                "page_number": idx,
                "section_name": f"Incident Record {rec_id}",
                "machine_type": m_type,
                "component": failure_mode,
                "uploaded_at": datetime.utcnow().isoformat(),
                "content_hash": compute_text_hash(semantic_record)
            }
            pages.append(ParsedPage(semantic_record, idx, f"Incident {rec_id}", meta))

    return ParsedDocument(document_id, document_name, "CSV", category, pages, file_hash)

def parse_document(file_path: Path, category: str) -> ParsedDocument:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(file_path, category)
    elif suffix == ".docx":
        return parse_docx(file_path, category)
    elif suffix == ".txt":
        return parse_txt(file_path, category)
    elif suffix == ".csv":
        return parse_csv(file_path, category)
    else:
        raise ValueError(f"Unsupported file format: {suffix}")
